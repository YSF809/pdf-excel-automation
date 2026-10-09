from pathlib import Path
from datetime import datetime
import re
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


try:
    import pdfplumber
except ModuleNotFoundError:
    raise SystemExit(
        "Missing dependency: 'pdfplumber'. Install it with: pip install pdfplumber"
    )

import pandas as pd


INPUT_FOLDER = Path("input_files")
OUTPUT_FOLDER = Path("output")



def find_value(pattern, text):
    match = re.search(
        pattern,
        text,
        re.IGNORECASE | re.MULTILINE
    )
    return match.group(1).strip() if match else ""


def extract_amount(label_pattern, text):
    pattern = (
        rf"^[ \t]*{label_pattern}[ \t]*:[ \t]*"
        r"(?:\$\s*|SAR[ \t]*)?"
        r"([0-9][0-9,]*(?:\.[0-9]+)?)"
        r"(?:[ \t]*SAR)?[ \t]*$"
    )

    value = find_value(pattern, text)

    return float(value.replace(",", "")) if value else None


def extract_pdf_data(file_path):
    print(f"Reading PDF: {file_path.name}")

    with pdfplumber.open(file_path) as pdf:
        text = "\n".join(
            page.extract_text() or ""
            for page in pdf.pages
        )

    invoice_number = find_value(
        r"^[ \t]*(INV-[A-Za-z0-9-]+)[ \t]*$",
        text
    )

    if not invoice_number:
        invoice_number = find_value(
            r"^[ \t]*Invoice[ \t]*(?:#|No\.?|Number)"
            r"[ \t]*:[ \t]*([A-Za-z0-9-]+)",
            text
        )

    date = find_value(
        r"^[ \t]*(?:Date[ \t]*:[ \t]*)?"
        r"(\d{4}-\d{2}-\d{2})[ \t]*$",
        text
    )

    customer = find_value(
        r"^Bill[ \t]+To[ \t]*:[ \t]*\r?\n"
        r"[ \t]*([^\r\n]+)",
        text
    )

    subtotal = extract_amount("Subtotal", text)

    vat = extract_amount(
        r"VAT(?:[ \t]*\([^\)\r\n]+\))?",
        text
    )

    total = extract_amount("TOTAL", text)

    complete = (
        bool(invoice_number and date and customer)
        and all(
            value is not None
            for value in (subtotal, vat, total)
        )
    )

    return {
        "Source File": file_path.name,
        "Invoice Number": invoice_number,
        "Date": date,
        "Customer": customer,
        "Subtotal": subtotal,
        "VAT": vat,
        "Total": total,
        "Status": "Extracted" if complete else "Needs Review",
        "Raw Text": text[:30000]
    }



def clean_csv(file_path):

    print(f"Cleaning CSV: {file_path.name}")

    try:
        df = pd.read_csv(file_path)

    except UnicodeDecodeError:
        df = pd.read_csv(
            file_path,
            encoding="latin-1"
        )

    df.columns = df.columns.str.strip()

    for column in df.select_dtypes(
        include="object"
    ).columns:

        df[column] = df[column].map(
            lambda value: (
                value.strip()
                if isinstance(value, str)
                else value
            )
        )

    df = df.dropna(how="all")

    df = df.drop_duplicates()

    df.insert(
        0,
        "Source File",
        file_path.name
    )

    return df


def format_excel(file_path):
    workbook = load_workbook(file_path)

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 30

        headers = {
            cell.value: cell.column
            for cell in sheet[1]
        }

        widths = {
            "Source File": 30,
            "Invoice Number": 27,
            "Date": 17,
            "Customer": 25,
            "Subtotal": 19,
            "VAT": 17,
            "Total": 19,
            "Status": 20,
            "Raw Text": 45
        }

        for cell in sheet[1]:
            cell.fill = PatternFill(
                "solid",
                fgColor="17365D"
            )
            cell.font = Font(
                bold=True,
                color="FFFFFF",
                size=11
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

            letter = get_column_letter(cell.column)
            sheet.column_dimensions[letter].width = widths.get(
                cell.value, 22
            )

        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="center")

                if cell.row % 2 == 0:
                    cell.fill = PatternFill(
                        "solid",
                        fgColor="EDF3F8"
                    )

        for name in ["Subtotal", "VAT", "Total"]:
            if name in headers:
                col = headers[name]
                for row in range(2, sheet.max_row + 1):
                    sheet.cell(row, col).number_format = (
                        '#,##0.00 "SAR"'
                    )

        if "Status" in headers:
            col = headers["Status"]

            for row in range(2, sheet.max_row + 1):
                cell = sheet.cell(row, col)

                if cell.value == "Extracted":
                    cell.font = Font(
                        color="008000",
                        bold=True
                    )
                elif cell.value == "Needs Review":
                    cell.font = Font(
                        color="C00000",
                        bold=True
                    )

        if "Raw Text" in headers:
            letter = get_column_letter(headers["Raw Text"])
            sheet.column_dimensions[letter].hidden = True

    workbook.save(file_path)



def generate_excel():
    INPUT_FOLDER.mkdir(exist_ok=True)
    OUTPUT_FOLDER.mkdir(exist_ok=True)

    invoices = []
    csv_data = []
    errors = []

    pdf_files = sorted(INPUT_FOLDER.glob("*.pdf"))
    csv_files = sorted(INPUT_FOLDER.glob("*.csv"))

    if not pdf_files and not csv_files:
        raise RuntimeError("No PDF or CSV files were found.")

    for pdf_file in pdf_files:
        try:
            data = extract_pdf_data(pdf_file)
            invoices.append(data)
            print(f"[SUCCESS] {pdf_file.name}")

        except Exception as error:
            errors.append({
                "Source File": pdf_file.name,
                "Error": str(error)
            })
            print(f"[ERROR] {pdf_file.name}: {error}")

    for csv_file in csv_files:
        try:
            data = clean_csv(csv_file)
            csv_data.append(data)
            print(f"[SUCCESS] {csv_file.name}")

        except Exception as error:
            errors.append({
                "Source File": csv_file.name,
                "Error": str(error)
            })
            print(f"[ERROR] {csv_file.name}: {error}")

    if not invoices and not csv_data:
        details = "\n".join(
            f"{item['Source File']}: {item['Error']}"
            for item in errors
        )

        raise RuntimeError(
            "All files failed to process:\n\n" + details
        )

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    output_file = (
        OUTPUT_FOLDER / f"report_{timestamp}.xlsx"
    )

    with pd.ExcelWriter(
        output_file,
        engine="openpyxl"
    ) as writer:

        if invoices:
            pd.DataFrame(invoices).to_excel(
                writer,
                sheet_name="PDF Invoices",
                index=False
            )

        if csv_data:
            pd.concat(
                csv_data,
                ignore_index=True
            ).to_excel(
                writer,
                sheet_name="CSV Data",
                index=False
            )

        if errors:
            pd.DataFrame(errors).to_excel(
                writer,
                sheet_name="Processing Errors",
                index=False
            )

    format_excel(output_file)

    print("\n" + "=" * 45)
    print("REPORT GENERATED")
    print("=" * 45)
    print(f"PDFs processed: {len(pdf_files)}")
    print(f"CSVs processed: {len(csv_files)}")
    print(f"Errors: {len(errors)}")
    print(f"Output: {output_file}")

    return output_file
